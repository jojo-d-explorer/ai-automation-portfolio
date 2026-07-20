#!/usr/bin/env python3
"""
score.py — v4 Phase 2, judgment layer.

No existing script to refactor from — new build per
HANDOFF_ROADMAP_v4_Claude_Code.md §5. Reads package.py's deduped final
output (Verified_Jobs_LATAM_Joey_{date}.csv), batches JD text (5 per call,
per roadmap spec) through the locked prompts/score_jd_v4.md against
config/joey_resume_digest.md, and writes scores back.

Backend is a config option per roadmap §8 ("Resolved 2026-07-19"):
  --backend claude-p   (default) headless `claude -p`, for interactive
                        Claude Code sessions — no API cost.
  --backend api         Anthropic API (Sonnet) — "approved fallback for
                        scheduled/unattended runs... cost is not a blocker."

Anti-hallucination invariant (roadmap §5 item 3): only ever scores rows
with jd_fetch_status=success. This is already guaranteed by package.py's
input (it only carries success rows forward), but re-checked here
defensively rather than trusted blindly.

Usage:
    python3 JC3/score.py                    # claude-p backend, this week
    python3 JC3/score.py --backend api       # Anthropic API backend
"""

import csv
import hashlib
import json
import os
import re
import subprocess
import sys
import time
from datetime import datetime, timedelta

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
BASE_PATH = os.path.dirname(SCRIPT_DIR)
PROMPT_PATH = os.path.join(BASE_PATH, "prompts", "score_jd_v4.md")
RESUME_DIGEST_PATH = os.path.join(BASE_PATH, "config", "joey_resume_digest.md")

BATCH_SIZE = 5
API_MODEL = "claude-sonnet-5"

# Tool-forced structured output for the API backend — a freeform-text
# JSON array (score_jd_v4.md's literal contract, needed for the claude-p
# CLI backend below) turned out unreliable in practice: a live test run
# hit both markdown code-fencing despite "no code fences" in the prompt,
# and a genuine JSON syntax error mid-response (an unescaped character in
# a long rationale string broke strict parsing). Tool calling makes
# malformed output structurally impossible for the API path — same fix
# pattern as the grading-contract work elsewhere in this session. The
# claude-p CLI backend has no equivalent mechanism from a prompt string
# alone, so it keeps the freeform-text + _extract_json_array fallback.
SCORE_TOOL = {
    "name": "submit_scores",
    "description": "Submit scoring results for a batch of jobs, per prompts/score_jd_v4.md.",
    "input_schema": {
        "type": "object",
        "properties": {
            "scores": {
                "type": "array",
                "items": {
                    "type": "object",
                    "properties": {
                        "job_id": {"type": "string"},
                        "error": {"type": ["string", "null"], "description": "'insufficient_jd_text' if jd_text was missing/empty/under 500 chars; omit otherwise."},
                        "score": {"type": "number"},
                        "capped": {"type": "boolean"},
                        "disqualified": {"type": ["string", "null"]},
                        "subscores": {
                            "type": "object",
                            "properties": {
                                "skills": {"type": "number"}, "geography": {"type": "number"},
                                "seniority": {"type": "number"}, "sector": {"type": "number"},
                            },
                        },
                        "stretch_dimensions": {"type": "array", "items": {"type": "string"}},
                        "red_flags": {"type": "array", "items": {"type": "string"}},
                        "spanish_requirement": {"type": "string", "enum": ["required", "preferred", "not_mentioned"]},
                        "location_type": {"type": "string"},
                        "salary_usd": {"type": "string"},
                        "comp_band_type": {"type": "string", "enum": ["us_level", "latam_local", "unclear"]},
                        "reports_to": {"type": "string"},
                        "builder_vs_manager": {"type": "string", "enum": ["zero_to_one", "scale_existing", "maintain", "unclear"]},
                        "rationale": {"type": "string"},
                        "recommend": {"type": "string", "enum": ["apply", "review", "skip"]},
                    },
                    "required": ["job_id"],
                },
            },
        },
        "required": ["scores"],
    },
}

sys.path.insert(0, SCRIPT_DIR)
from package import _autosize_and_style, HEADER_FILL, HEADER_FONT  # noqa: E402


def _url_hash(url):
    return hashlib.sha256(url.encode("utf-8")).hexdigest()[:16]


def _week_paths():
    today = datetime.now()
    monday = today - timedelta(days=today.weekday())
    week_date = monday.strftime("%Y-%m-%d")
    week_dir = os.path.join(BASE_PATH, "results", "joey", "LATAM", f"Week_of_{week_date}")
    return week_date, week_dir


def _load_csv(path):
    if not os.path.exists(path):
        return []
    with open(path, newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def load_scoreable_rows(verified_csv):
    """Anti-hallucination invariant: only rows with cached JD text, i.e.
    jd_fetch_status=success (re-checked here, not just trusted from
    upstream) get built into scoring batches."""
    rows = _load_csv(verified_csv)
    scoreable = []
    for row in rows:
        if row.get("jd_fetch_status") != "success":
            continue
        cache_path = row.get("jd_cache_path", "")
        if not cache_path:
            continue
        full_cache_path = os.path.join(BASE_PATH, cache_path)
        if not os.path.exists(full_cache_path):
            continue
        with open(full_cache_path, encoding="utf-8") as f:
            jd_text = f.read()
        row["_job_id"] = _url_hash(row.get("url", ""))
        row["_jd_text"] = jd_text
        scoreable.append(row)
    return scoreable


def _batches(items, size):
    for i in range(0, len(items), size):
        yield items[i:i + size]


def _build_batch_payload(batch):
    return [
        {
            "job_id": r["_job_id"],
            "company": r.get("company_name", ""),
            "title": r.get("title", ""),
            "source": r.get("found_via") or r.get("source", ""),
            "location_string": r.get("location", ""),
            "jd_text": r["_jd_text"],
        }
        for r in batch
    ]


def _extract_json_array(text):
    """Score_jd_v4.md's contract is "return ONLY the JSON array" — but a
    CLI wrapper (claude -p --output-format json) may envelope that in its
    own JSON with a result/content field. Try direct parse first, then
    common envelope keys, then a regex fallback for a bare [...] block."""
    text = text.strip()
    try:
        parsed = json.loads(text)
        if isinstance(parsed, list):
            return parsed
        if isinstance(parsed, dict):
            for key in ("result", "content", "text", "output"):
                if key in parsed:
                    inner = parsed[key]
                    if isinstance(inner, list):
                        return inner
                    if isinstance(inner, str):
                        return _extract_json_array(inner)
    except json.JSONDecodeError:
        pass

    match = re.search(r"\[.*\]", text, re.DOTALL)
    if match:
        return json.loads(match.group(0))

    raise ValueError(f"Could not extract a JSON array from model output: {text[:200]}...")


def score_batch_claude_p(system_prompt, user_payload):
    full_prompt = f"{system_prompt}\n\n---\n\nJobs to score:\n\n{json.dumps(user_payload, ensure_ascii=False)}"
    result = subprocess.run(
        ["claude", "-p", full_prompt, "--output-format", "json"],
        capture_output=True, text=True, timeout=120,
    )
    if result.returncode != 0:
        raise RuntimeError(f"claude -p failed (exit {result.returncode}): {result.stderr[:300]}")
    return _extract_json_array(result.stdout)


def score_batch_api(system_prompt, user_payload, client):
    message = client.messages.create(
        model=API_MODEL,
        max_tokens=4096,
        system=system_prompt,
        tools=[SCORE_TOOL],
        tool_choice={"type": "tool", "name": "submit_scores"},
        messages=[{
            "role": "user",
            "content": f"Jobs to score:\n\n{json.dumps(user_payload, ensure_ascii=False)}",
        }],
    )
    tool_use = next((b for b in message.content if b.type == "tool_use"), None)
    if tool_use is None:
        raise RuntimeError("API response had no tool_use block")
    return tool_use.input["scores"]


def run_score(verified_csv, backend="claude-p"):
    with open(PROMPT_PATH, encoding="utf-8") as f:
        prompt_body = f.read()
    with open(RESUME_DIGEST_PATH, encoding="utf-8") as f:
        resume_digest = f.read()
    system_prompt = f"{prompt_body}\n\n---\n\nResume digest:\n\n{resume_digest}"

    scoreable = load_scoreable_rows(verified_csv)
    print(f"\n{'─'*70}")
    print(f"  score.py — {len(scoreable)} scoreable rows (backend={backend})")
    print(f"{'─'*70}")

    client = None
    if backend == "api":
        import anthropic
        client = anthropic.Anthropic()  # reads ANTHROPIC_API_KEY from env

    scores_by_job_id = {}
    errors = []

    for batch_num, batch in enumerate(_batches(scoreable, BATCH_SIZE), 1):
        payload = _build_batch_payload(batch)
        print(f"  batch {batch_num} ({len(batch)} jobs)...")
        try:
            if backend == "claude-p":
                results = score_batch_claude_p(system_prompt, payload)
            else:
                results = score_batch_api(system_prompt, payload, client)
        except Exception as e:
            print(f"    ERROR: {e}")
            errors.extend(r["_job_id"] for r in batch)
            continue

        for result in results:
            job_id = result.get("job_id")
            if not job_id:
                continue
            if result.get("error"):
                print(f"    {job_id}: {result['error']}")
                errors.append(job_id)
                continue
            scores_by_job_id[job_id] = result
        time.sleep(0.5)

    print(f"\n  Scored: {len(scores_by_job_id)} | Errors/skipped: {len(errors)}")

    for row in scoreable:
        result = scores_by_job_id.get(row["_job_id"])
        if not result:
            continue
        row["score"] = result.get("score", "")
        row["capped"] = result.get("capped", "")
        row["disqualified"] = result.get("disqualified", "") or ""
        row["recommend"] = result.get("recommend", "")
        row["rationale"] = result.get("rationale", "")
        row["red_flags"] = ", ".join(result.get("red_flags", []) or [])
        row["stretch_dimensions"] = ", ".join(result.get("stretch_dimensions", []) or [])
        row["spanish_requirement"] = result.get("spanish_requirement", "")
        row["location_type"] = result.get("location_type", "")
        row["salary_usd"] = result.get("salary_usd", "")
        row["comp_band_type"] = result.get("comp_band_type", "")
        row["reports_to"] = result.get("reports_to", "")
        row["builder_vs_manager"] = result.get("builder_vs_manager", "")
        subscores = result.get("subscores", {}) or {}
        row["subscore_skills"] = subscores.get("skills", "")
        row["subscore_geography"] = subscores.get("geography", "")
        row["subscore_seniority"] = subscores.get("seniority", "")
        row["subscore_sector"] = subscores.get("sector", "")

    scored_rows = [r for r in scoreable if r["_job_id"] in scores_by_job_id]
    for r in scored_rows:
        del r["_job_id"]
        del r["_jd_text"]

    def _sort_key(r):
        try:
            return -float(r.get("score") or 0)
        except (TypeError, ValueError):
            return 0
    scored_rows.sort(key=_sort_key)

    return scored_rows


def write_outputs(scored_rows, week_dir, week_date):
    csv_path = os.path.join(week_dir, f"Scored_Jobs_LATAM_Joey_{week_date}.csv")
    if scored_rows:
        fieldnames = list(scored_rows[0].keys())
        with open(csv_path, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=fieldnames)
            writer.writeheader()
            writer.writerows(scored_rows)

    from openpyxl import Workbook
    xlsx_path = os.path.join(week_dir, f"Scored_LATAM_Joey_Complete_{week_date}.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Scored Jobs"
    _autosize_and_style(ws, scored_rows)
    if scored_rows:
        recommend_col = list(scored_rows[0].keys()).index("recommend") + 1
        from openpyxl.styles import PatternFill
        apply_fill = PatternFill(start_color="C8E6C9", end_color="C8E6C9", fill_type="solid")
        review_fill = PatternFill(start_color="FFF9C4", end_color="FFF9C4", fill_type="solid")
        for row_idx, row in enumerate(scored_rows, start=2):
            if row.get("recommend") == "apply":
                ws.cell(row=row_idx, column=recommend_col).fill = apply_fill
            elif row.get("recommend") == "review":
                ws.cell(row=row_idx, column=recommend_col).fill = review_fill
    wb.save(xlsx_path)

    return csv_path, xlsx_path


def main():
    backend = "claude-p"
    if "--backend" in sys.argv:
        idx = sys.argv.index("--backend")
        backend = sys.argv[idx + 1]
    if backend not in ("claude-p", "api"):
        print(f"ERROR: --backend must be 'claude-p' or 'api', got {backend!r}")
        sys.exit(1)

    week_date, week_dir = _week_paths()
    verified_csv = os.path.join(week_dir, f"Verified_Jobs_LATAM_Joey_{week_date}.csv")
    if not os.path.exists(verified_csv):
        print(f"No verified jobs CSV found at {verified_csv}. Run the full pipeline "
              f"(discover -> health -> fetch_jds -> package) first.")
        sys.exit(1)

    scored_rows = run_score(verified_csv, backend=backend)
    csv_path, xlsx_path = write_outputs(scored_rows, week_dir, week_date)

    apply_n = sum(1 for r in scored_rows if r.get("recommend") == "apply")
    review_n = sum(1 for r in scored_rows if r.get("recommend") == "review")
    skip_n = sum(1 for r in scored_rows if r.get("recommend") == "skip")

    print(f"\n{'='*70}")
    print(f"  Scored {len(scored_rows)} jobs — apply: {apply_n}, review: {review_n}, skip: {skip_n}")
    print(f"  CSV:  {csv_path}")
    print(f"  XLSX: {xlsx_path}")
    print(f"{'='*70}\n")


if __name__ == "__main__":
    main()
